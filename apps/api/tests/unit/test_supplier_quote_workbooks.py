from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook

from app.supplier_quotes.schemas import WorkbookMapping
from app.supplier_quotes.workbooks import inspect_supplier_workbook, parse_supplier_workbook


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    cover = workbook.active
    cover.title = "Cover"
    cover.append(["Supplier price list"])

    offers = workbook.create_sheet("Offers")
    offers.append(["Updated price list"])
    offers.append(["Prepared for purchasing"])
    offers.append(["Item code", "Product", "Price USD", "MOQ", "Unit"])
    offers.append(["AX-1", "Widget", 12.3, None, "piece"])
    offers.append(["AX-2", "Gadget", 8, 10, "box"])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _quoted_workbook_bytes(*, price: object = 12.3, repeated_header: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Offers"
    sheet.append(["SKU", "Name", "Price", "MOQ", "Unit"])
    sheet.append(["A-1", "Widget", price, None, "piece"])
    if repeated_header:
        sheet.append(["SKU", "Name", "Price", "MOQ", "Unit"])
    sheet.append(["A-2", "Gadget", 8, 5, "piece"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _mapping(*, worksheet_name: str = "Offers") -> WorkbookMapping:
    return WorkbookMapping(
        worksheet_name=worksheet_name,
        header_row=1,
        columns={
            "supplier_sku": "SKU",
            "product_name": "Name",
            "unit_price": "Price",
            "minimum_order_quantity": "MOQ",
            "unit": "Unit",
        },
        default_currency="USD",
    )


def test_inspection_finds_headers_after_title_rows_across_sheets() -> None:
    inspection = inspect_supplier_workbook(_workbook_bytes())

    assert [sheet.name for sheet in inspection.sheets] == ["Cover", "Offers"]
    offers = inspection.sheets[1]
    assert 3 in offers.candidate_header_rows
    assert offers.header_labels[3] == ["Item code", "Product", "Price USD", "MOQ", "Unit"]


def test_inspection_caps_sample_rows_at_five_for_the_whole_workbook() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["SKU", "Name", "Price"])
    for number in range(20):
        sheet.append([f"SKU-{number}", f"Product {number}", number + 1])
    output = BytesIO()
    workbook.save(output)

    inspection = inspect_supplier_workbook(output.getvalue())

    assert sum(len(sheet.sample_rows) for sheet in inspection.sheets) <= 5


def test_parser_uses_confirmed_header_mapping_and_preserves_source_rows() -> None:
    mapping = _mapping(worksheet_name="Offers")
    mapping = mapping.model_copy(
        update={
            "header_row": 3,
            "columns": {
                "supplier_sku": "Item code",
                "product_name": "Product",
                "unit_price": "Price USD",
                "minimum_order_quantity": "MOQ",
                "unit": "Unit",
            },
        }
    )

    parsed = parse_supplier_workbook(_workbook_bytes(), mapping)

    assert len(parsed.rows) == 2
    first = parsed.rows[0]
    assert first.worksheet_name == "Offers"
    assert first.source_row_number == 4
    assert first.supplier_sku == "AX-1"
    assert first.product_name == "Widget"
    assert first.unit_price == Decimal("12.3")
    assert first.currency == "USD"
    assert first.purchase_unit == "piece"
    assert first.minimum_order_quantity is None
    assert parsed.rows[1].source_row_number == 5
    assert parsed.rows[1].minimum_order_quantity == Decimal("10")
    assert parsed.issues == []


def test_missing_worksheet_returns_a_bounded_issue() -> None:
    parsed = parse_supplier_workbook(_workbook_bytes(), _mapping(worksheet_name="No sheet"))

    assert parsed.rows == []
    assert [issue.code for issue in parsed.issues] == ["worksheet_not_found"]


def test_invalid_price_does_not_discard_other_valid_rows() -> None:
    parsed = parse_supplier_workbook(_quoted_workbook_bytes(price="call for price"), _mapping())

    assert [row.supplier_sku for row in parsed.rows] == ["A-2"]
    assert [(issue.source_row_number, issue.code) for issue in parsed.issues] == [
        (2, "invalid_price")
    ]


def test_repeated_header_rows_are_reported_and_skipped() -> None:
    parsed = parse_supplier_workbook(_quoted_workbook_bytes(repeated_header=True), _mapping())

    assert [row.source_row_number for row in parsed.rows] == [2, 4]
    assert [(issue.source_row_number, issue.code) for issue in parsed.issues] == [
        (3, "repeated_header")
    ]


def test_merged_cells_in_a_data_row_are_reported_without_reading_a_fake_price() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Offers"
    sheet.append(["SKU", "Name", "Price", "MOQ", "Unit"])
    sheet.append(["A-1", "Widget", 12.3, None, "piece"])
    sheet.merge_cells("B2:C2")
    output = BytesIO()
    workbook.save(output)

    parsed = parse_supplier_workbook(output.getvalue(), _mapping())

    assert parsed.rows == []
    assert [(issue.source_row_number, issue.code) for issue in parsed.issues] == [
        (2, "merged_data_cells")
    ]


def test_formula_without_cached_value_is_not_evaluated() -> None:
    parsed = parse_supplier_workbook(_quoted_workbook_bytes(price="=6*2"), _mapping())

    assert [row.supplier_sku for row in parsed.rows] == ["A-2"]
    assert [(issue.source_row_number, issue.code) for issue in parsed.issues] == [
        (2, "invalid_price")
    ]


def test_invalid_xlsx_payload_returns_a_bounded_issue() -> None:
    parsed = parse_supplier_workbook(b"not a zip workbook", _mapping())

    assert parsed.rows == []
    assert [issue.code for issue in parsed.issues] == ["invalid_workbook"]
    assert len(parsed.issues[0].message) <= 160


def test_oversized_attachment_is_rejected_before_workbook_loading() -> None:
    parsed = parse_supplier_workbook(b"x" * (20 * 1024 * 1024 + 1), _mapping())

    assert parsed.rows == []
    assert [issue.code for issue in parsed.issues] == ["file_too_large"]


def test_worksheet_over_row_limit_is_rejected() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Offers"
    sheet.append(["SKU", "Name", "Price", "MOQ", "Unit"])
    sheet["A25001"] = "out-of-range"
    output = BytesIO()
    workbook.save(output)

    parsed = parse_supplier_workbook(output.getvalue(), _mapping())

    assert parsed.rows == []
    assert [issue.code for issue in parsed.issues] == ["row_limit_exceeded"]


def _with_broken_sheet_xml(data: bytes) -> bytes:
    import zipfile

    source = zipfile.ZipFile(BytesIO(data))
    output = BytesIO()
    with zipfile.ZipFile(output, "w") as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename.startswith("xl/worksheets/sheet"):
                content = content[: len(content) // 2]
            target.writestr(entry, content)
    return output.getvalue()


def test_malformed_sheet_xml_returns_a_bounded_issue_instead_of_raising() -> None:
    data = _with_broken_sheet_xml(_quoted_workbook_bytes())

    inspection = inspect_supplier_workbook(data)
    result = parse_supplier_workbook(data, _mapping())

    assert [issue.code for issue in inspection.issues] == ["invalid_workbook"]
    assert result.rows == []
    assert [issue.code for issue in result.issues] == ["invalid_workbook"]
