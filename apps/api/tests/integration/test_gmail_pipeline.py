"""End-to-end Gmail supplier quote pipeline against the real database.

Google is faked at the `GmailClient` / `GmailOAuthClient` boundary and the
mailbox lock at Redis; everything between -- history paging, per-message
processing, workbook parsing, quote persistence, reports and the history
cursor -- is the production code.
"""

from __future__ import annotations

import base64
import uuid
from collections.abc import Callable
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from openpyxl import Workbook
from sqlalchemy import select, text

from app.core.ids import uuid7
from app.core.tenancy import TenantContext, tenant_session
from app.db.models import GmailConnection, Product, SupplierMessage, SupplierReport
from app.gmail.client import GmailAPIError, GmailClient
from app.gmail.oauth import GmailOAuthClient, OAuthTokens
from app.supplier_quotes.schemas import SaveSupplierMappingInput, SupplierInput, WorkbookMapping
from app.supplier_quotes.service import SupplierQuoteService

pytestmark = pytest.mark.anyio

_KEY = b"p" * 32
_SENDER = "quotes@acme.example"
_MAILBOX = "buyer@example.com"


def workbook_bytes(rows: list[tuple[Any, ...]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Offers"
    sheet.append(("SKU", "Name", "Price"))
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def gmail_message(
    message_id: str,
    *,
    attachment_id: str = "att-1",
    part_id: str = "1",
    size: int = 4096,
    sender: str = _SENDER,
) -> dict[str, Any]:
    return {
        "id": message_id,
        "labelIds": ["INBOX"],
        "internalDate": "1790000000000",
        "snippet": "Please find our prices attached.",
        "payload": {
            "headers": [
                {"name": "From", "value": f"Acme <{sender}>"},
                {"name": "Subject", "value": "Price list"},
            ],
            "parts": [
                {"partId": "0", "mimeType": "text/plain", "body": {"size": 10}},
                {
                    "partId": part_id,
                    "filename": "prices.xlsx",
                    "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    "body": {"attachmentId": attachment_id, "size": size},
                },
            ],
        },
    }


class FakeGmail:
    """Messages, attachments and history served to the production client calls."""

    def __init__(self) -> None:
        self.messages: dict[str, dict[str, Any]] = {}
        self.attachments: dict[str, bytes] = {}
        self.history_message_ids: list[str] = []
        self.history_id = "200"
        self.downloaded: list[str] = []
        self.fail_message: Callable[[str], None] | None = None

    def install(self, monkeypatch: pytest.MonkeyPatch) -> None:
        fake = self

        async def list_history(_self: GmailClient, _token: str, _cursor: str, **_k: object) -> dict:
            return {
                "history": [
                    {"id": fake.history_id, "messagesAdded": [{"message": {"id": mid}}]}
                    for mid in fake.history_message_ids
                ],
                "historyId": fake.history_id,
            }

        async def get_message(_self: GmailClient, _token: str, message_id: str) -> dict:
            if fake.fail_message is not None:
                fake.fail_message(message_id)
            return fake.messages[message_id]

        async def get_attachment(
            _self: GmailClient, _token: str, _message_id: str, attachment_id: str
        ) -> bytes:
            fake.downloaded.append(attachment_id)
            return fake.attachments[attachment_id]

        async def refresh(_self: GmailOAuthClient, _token: str, **_k: object) -> OAuthTokens:
            return OAuthTokens(access_token="access-token")

        monkeypatch.setattr(GmailClient, "list_history", list_history)
        monkeypatch.setattr(GmailClient, "get_message", get_message)
        monkeypatch.setattr(GmailClient, "get_attachment", get_attachment)
        monkeypatch.setattr(GmailOAuthClient, "refresh_access_token", refresh)


class _Lock:
    def __init__(self, held: bool = False) -> None:
        self.held = held

    async def acquire(self, *, blocking: bool) -> bool:
        return not self.held

    async def release(self) -> None:
        return None


class _Redis:
    def __init__(self, held: bool = False) -> None:
        self.held = held

    def lock(self, *_args: object, **_kwargs: object) -> _Lock:
        return _Lock(self.held)


@pytest.fixture
def fake_gmail(monkeypatch: pytest.MonkeyPatch) -> FakeGmail:
    from app.core.config import get_settings

    settings = get_settings().model_copy(
        update={"gmail_token_encryption_key": base64.b64encode(_KEY).decode()}
    )
    monkeypatch.setattr("app.gmail.token_crypto.get_settings", lambda: settings)
    monkeypatch.setattr("app.gmail.sync.get_redis", lambda: _Redis())
    fake = FakeGmail()
    fake.install(monkeypatch)
    return fake


async def seed_mailbox(
    owner_connection: Any, tenant: TenantContext, *, confirmed_layout: bool = True
) -> uuid.UUID:
    """A connected mailbox, a supplier (with a confirmed layout unless told
    otherwise), and a priced product."""
    from app.gmail.token_crypto import encrypt_token

    connection_id = uuid7()
    await owner_connection.execute(
        text(
            "INSERT INTO gmail_connections "
            "(id, organization_id, email, refresh_token_ciphertext, history_id, status) "
            "VALUES (:id, :org, :email, :token, '100', 'connected')"
        ),
        {
            "id": connection_id,
            "org": tenant.organization_id,
            "email": _MAILBOX,
            "token": encrypt_token("refresh-token", key=_KEY),
        },
    )
    await owner_connection.commit()
    async with tenant_session(tenant) as session:
        service = SupplierQuoteService(session, tenant)
        supplier = await service.create_supplier(
            SupplierInput(name="Acme", sender_emails=[_SENDER])
        )
        if confirmed_layout:
            await service.save_mapping(
                supplier.id,
                SaveSupplierMappingInput(
                    mapping=WorkbookMapping(
                        worksheet_name="Offers",
                        header_row=1,
                        columns={
                            "supplier_sku": "SKU",
                            "product_name": "Name",
                            "unit_price": "Price",
                        },
                        default_currency="USD",
                        default_purchase_unit="each",
                    ),
                    column_labels=["SKU", "Name", "Price"],
                    is_confirmed=True,
                ),
            )
        session.add(
            Product(
                organization_id=tenant.organization_id,
                external_id="SKU-1",
                name="Widget",
                slug="widget",
                price=Decimal("15.00"),
                currency="USD",
                attributes={"sale_unit": "each"},
            )
        )
    return connection_id


async def message_rows(tenant: TenantContext) -> dict[str, SupplierMessage]:
    async with tenant_session(tenant) as session:
        result = await session.execute(select(SupplierMessage))
        return {row.gmail_message_id: row for row in result.scalars().all()}


async def connection_row(tenant: TenantContext) -> GmailConnection:
    async with tenant_session(tenant) as session:
        result = await session.execute(select(GmailConnection))
        return result.scalar_one()


async def report_for(tenant: TenantContext, message: SupplierMessage) -> SupplierReport | None:
    async with tenant_session(tenant) as session:
        result = await session.execute(
            select(SupplierReport).where(SupplierReport.supplier_message_id == message.id)
        )
        return result.scalar_one_or_none()


async def test_a_message_that_fails_is_recorded_and_the_mailbox_keeps_moving(
    owner_connection: Any,
    tenant_a: TenantContext,
    fake_gmail: FakeGmail,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import app.gmail.sync as sync_module
    from app.gmail.sync import run_gmail_sync

    await seed_mailbox(owner_connection, tenant_a)
    fake_gmail.messages["m-bad"] = gmail_message("m-bad", attachment_id="att-bad")
    fake_gmail.messages["m-good"] = gmail_message("m-good", attachment_id="att-good")
    fake_gmail.attachments["att-bad"] = b"bad workbook"
    fake_gmail.attachments["att-good"] = workbook_bytes([("SKU-1", "Widget", 10)])
    fake_gmail.history_message_ids = ["m-bad", "m-good"]
    real_inspect = sync_module.inspect_supplier_workbook

    # Stands in for any failure nobody anticipated: the contract under test
    # is that one message cannot stop the mailbox, whatever it raises.
    def inspect(data: bytes) -> Any:
        if data == b"bad workbook":
            raise RuntimeError("unexpected parser failure")
        return real_inspect(data)

    monkeypatch.setattr(sync_module, "inspect_supplier_workbook", inspect)

    await run_gmail_sync(mailbox_email=_MAILBOX, notification_history_id="200")

    rows = await message_rows(tenant_a)
    assert rows["m-bad"].status == "failed"
    assert rows["m-bad"].error_code == "processing_failed"
    assert rows["m-good"].status == "processed"
    assert await report_for(tenant_a, rows["m-good"]) is not None
    connection = await connection_row(tenant_a)
    assert connection.history_id == "200"


async def test_a_transient_gmail_error_keeps_the_cursor_for_a_retry(
    owner_connection: Any, tenant_a: TenantContext, fake_gmail: FakeGmail
) -> None:
    from arq import Retry

    from app.gmail.sync import run_gmail_sync

    await seed_mailbox(owner_connection, tenant_a)
    fake_gmail.messages["m-1"] = gmail_message("m-1")
    fake_gmail.history_message_ids = ["m-1"]

    def unavailable(_message_id: str) -> None:
        raise GmailAPIError("Gmail is unavailable", status_code=503)

    fake_gmail.fail_message = unavailable

    with pytest.raises(Retry):
        await run_gmail_sync(mailbox_email=_MAILBOX, notification_history_id="200")

    assert (await connection_row(tenant_a)).history_id == "100"


async def test_an_oversized_attachment_is_reported_without_downloading_it(
    owner_connection: Any, tenant_a: TenantContext, fake_gmail: FakeGmail
) -> None:
    from app.gmail.sync import run_gmail_sync

    await seed_mailbox(owner_connection, tenant_a)
    fake_gmail.messages["m-1"] = gmail_message("m-1", size=25 * 1024 * 1024)
    fake_gmail.history_message_ids = ["m-1"]

    await run_gmail_sync(mailbox_email=_MAILBOX, notification_history_id="200")

    assert fake_gmail.downloaded == []
    message = (await message_rows(tenant_a))["m-1"]
    assert message.status == "review_required"
    report = await report_for(tenant_a, message)
    assert report is not None
    codes = [issue["code"] for issue in report.summary_json["row_issues"]]
    assert "file_too_large" in codes
    assert (await connection_row(tenant_a)).history_id == "200"


async def test_a_busy_mailbox_is_retried_instead_of_dropped(
    owner_connection: Any,
    tenant_a: TenantContext,
    fake_gmail: FakeGmail,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from arq import Retry

    from app.gmail.sync import run_gmail_sync

    await seed_mailbox(owner_connection, tenant_a)
    monkeypatch.setattr("app.gmail.sync.get_redis", lambda: _Redis(held=True))

    with pytest.raises(Retry):
        await run_gmail_sync(mailbox_email=_MAILBOX, notification_history_id="200")


async def test_a_busy_mailbox_is_left_to_the_next_poll(
    owner_connection: Any,
    tenant_a: TenantContext,
    fake_gmail: FakeGmail,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.gmail.sync import run_gmail_sync

    await seed_mailbox(owner_connection, tenant_a)
    fake_gmail.history_message_ids = ["m-busy"]
    monkeypatch.setattr("app.gmail.sync.get_redis", lambda: _Redis(held=True))

    # No Retry: the job holding the lock is syncing this mailbox right now,
    # and the next poll covers anything it listed too early to see.
    await run_gmail_sync(
        mailbox_email=_MAILBOX,
        notification_history_id=None,
        force_sync=True,
        skip_if_busy=True,
    )

    assert "m-busy" not in await message_rows(tenant_a)
    assert (await connection_row(tenant_a)).history_id == "100"


async def test_maintenance_continues_past_a_mailbox_that_fails(
    owner_connection: Any,
    tenant_a: TenantContext,
    tenant_b: TenantContext,
    fake_gmail: FakeGmail,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import app.gmail.sync as sync_module
    from app.gmail.client import GmailWatchResult

    await seed_mailbox(owner_connection, tenant_a)
    from app.gmail.token_crypto import encrypt_token

    await owner_connection.execute(
        text(
            "INSERT INTO gmail_connections "
            "(id, organization_id, email, refresh_token_ciphertext, history_id, status) "
            "VALUES (:id, :org, 'other@example.com', :token, '100', 'connected')"
        ),
        {
            "id": uuid7(),
            "org": tenant_b.organization_id,
            "token": encrypt_token("refresh-token", key=_KEY),
        },
    )
    await owner_connection.commit()
    synced: list[str] = []

    async def sync(*, mailbox_email: str, **_kwargs: object) -> None:
        synced.append(mailbox_email)
        if mailbox_email == _MAILBOX:
            raise RuntimeError("one tenant's mailbox is broken")

    async def watch(_self: GmailClient, _token: str, _topic: str) -> GmailWatchResult:
        return GmailWatchResult(history_id="100", expiration=datetime.now(UTC) + timedelta(days=7))

    monkeypatch.setattr(sync_module, "run_gmail_sync", sync)
    monkeypatch.setattr(GmailClient, "watch", watch)

    await sync_module.run_gmail_maintenance()

    # Other tests' organizations may leave mailboxes behind; what matters is
    # that the failing one did not stop the other from being maintained.
    assert {_MAILBOX, "other@example.com"} <= set(synced)
    assert (await connection_row(tenant_a)).last_error_code == "gmail_maintenance_failed"


async def test_a_confirmed_layout_reprocesses_even_when_gmail_reissues_attachment_ids(
    owner_connection: Any, tenant_a: TenantContext, fake_gmail: FakeGmail
) -> None:
    from app.gmail.sync import run_gmail_message_reprocess, run_gmail_sync
    from app.supplier_quotes.schemas import WorkbookMapping

    await seed_mailbox(owner_connection, tenant_a, confirmed_layout=False)
    fake_gmail.messages["m-1"] = gmail_message("m-1", attachment_id="att-first-fetch")
    fake_gmail.attachments["att-first-fetch"] = workbook_bytes([("SKU-1", "Widget", 10)])
    fake_gmail.history_message_ids = ["m-1"]

    await run_gmail_sync(mailbox_email=_MAILBOX, notification_history_id="200")

    message = (await message_rows(tenant_a))["m-1"]
    assert message.status == "review_required"
    async with tenant_session(tenant_a) as session:
        service = SupplierQuoteService(session, tenant_a)
        [proposal] = await service.list_pending_mappings()
        confirmed = await service.confirm_mapping(
            proposal.id,
            mapping=WorkbookMapping(
                worksheet_name="Offers",
                header_row=1,
                columns={"supplier_sku": "SKU", "product_name": "Name", "unit_price": "Price"},
                default_currency="USD",
                default_purchase_unit="each",
            ),
        )
        mapping_id = confirmed.id

    # Gmail hands out a fresh attachmentId on the next messages.get.
    fake_gmail.messages["m-1"] = gmail_message("m-1", attachment_id="att-second-fetch")
    fake_gmail.attachments["att-second-fetch"] = fake_gmail.attachments.pop("att-first-fetch")

    await run_gmail_message_reprocess(
        organization_id=str(tenant_a.organization_id),
        supplier_message_id=str(message.id),
        mapping_id=str(mapping_id),
    )

    message = (await message_rows(tenant_a))["m-1"]
    assert message.status == "processed"
    async with tenant_session(tenant_a) as session:
        report = await report_for(tenant_a, message)
        assert report is not None
        _report, quotes = await SupplierQuoteService(session, tenant_a).get_report(report.id)
    assert [(quote.supplier_sku, quote.match_status) for quote in quotes] == [
        ("SKU-1", "confirmed")
    ]


async def test_a_new_layout_is_proposed_on_the_first_sheet_that_has_headers(
    owner_connection: Any, tenant_a: TenantContext, fake_gmail: FakeGmail
) -> None:
    from app.gmail.sync import run_gmail_sync

    workbook = Workbook()
    cover = workbook.active
    assert cover is not None
    cover.title = "Cover"
    cover.append(("Thank you for your enquiry",))
    offers = workbook.create_sheet("Offers")
    offers.append(("SKU", "Name", "Price"))
    offers.append(("SKU-1", "Widget", 10))
    buffer = BytesIO()
    workbook.save(buffer)
    await seed_mailbox(owner_connection, tenant_a, confirmed_layout=False)
    fake_gmail.messages["m-1"] = gmail_message("m-1")
    fake_gmail.attachments["att-1"] = buffer.getvalue()
    fake_gmail.history_message_ids = ["m-1"]

    await run_gmail_sync(mailbox_email=_MAILBOX, notification_history_id="200")

    async with tenant_session(tenant_a) as session:
        [proposal] = await SupplierQuoteService(session, tenant_a).list_pending_mappings()
    assert (proposal.worksheet_name, proposal.header_row) == ("Offers", 1)
    assert {"worksheet_name": "Offers", "header_row": 1, "labels": ["SKU", "Name", "Price"]} in (
        proposal.sheet_options
    )


async def test_a_remembered_supplier_sku_matches_without_review(
    owner_connection: Any, tenant_a: TenantContext, fake_gmail: FakeGmail
) -> None:
    from app.db.models import Supplier
    from app.gmail.sync import run_gmail_sync

    await seed_mailbox(owner_connection, tenant_a)
    async with tenant_session(tenant_a) as session:
        product = await session.scalar(select(Product).where(Product.external_id == "SKU-1"))
        supplier = await session.scalar(select(Supplier))
        assert product is not None and supplier is not None
        supplier.sku_aliases = {"ACME-77": str(product.id)}
    fake_gmail.messages["m-1"] = gmail_message("m-1")
    # The supplier's own SKU, not the catalogue's: only the alias links them.
    fake_gmail.attachments["att-1"] = workbook_bytes([("ACME-77", "Gadget", 10)])
    fake_gmail.history_message_ids = ["m-1"]

    await run_gmail_sync(mailbox_email=_MAILBOX, notification_history_id="200")

    message = (await message_rows(tenant_a))["m-1"]
    assert message.status == "processed"
    async with tenant_session(tenant_a) as session:
        report = await report_for(tenant_a, message)
        assert report is not None
        _report, [quote] = await SupplierQuoteService(session, tenant_a).get_report(report.id)
    assert (quote.match_status, quote.product_id) == ("confirmed", product.id)
