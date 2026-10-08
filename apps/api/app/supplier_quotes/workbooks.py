from __future__ import annotations

import posixpath
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from app.supplier_quotes.schemas import (
    CanonicalQuoteRow,
    QuantityTier,
    WorkbookField,
    WorkbookInspection,
    WorkbookMapping,
    WorkbookParseResult,
    WorkbookRowIssue,
    WorkbookSampleRow,
    WorkbookSheetInspection,
)

MAX_ATTACHMENT_BYTES = 20 * 1024 * 1024
MAX_WORKSHEET_ROWS = 25_000
MAX_PREVIEW_ROWS = 5
MAX_HEADER_SCAN_ROWS = 50
MAX_COLUMN_COUNT = 100
MAX_CELL_PREVIEW_CHARS = 160
MAX_ZIP_EXPANDED_BYTES = 100 * 1024 * 1024

_WORKBOOK_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_QUANTITY_TIER_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*[:=@]\s*(\d+(?:\.\d*)?|\.\d+)\s*$")


def inspect_supplier_workbook(data: bytes) -> WorkbookInspection:
    """Return a bounded, in-memory preview suitable for suggesting a column map."""
    workbook, issue = _open_workbook(data)
    if issue is not None:
        return WorkbookInspection(sheets=[], issues=[issue])
    assert workbook is not None

    remaining_samples = MAX_PREVIEW_ROWS
    inspected_sheets: list[WorkbookSheetInspection] = []
    try:
        for worksheet in workbook.worksheets:
            rows = _read_preview_rows(worksheet)
            labels: dict[int, list[str | None]] = {}
            candidates: list[int] = []
            for row_number, values in enumerate(rows, start=1):
                nonempty = [value for value in values if _has_value(value)]
                if len(nonempty) < 2:
                    continue
                string_count = sum(isinstance(value, str) for value in nonempty)
                if string_count / len(nonempty) < 0.5:
                    continue
                candidates.append(row_number)
                labels[row_number] = [
                    _preview_text(value) for value in _trim_trailing_empty(values)
                ]

            samples: list[WorkbookSampleRow] = []
            if remaining_samples and candidates:
                first_header = candidates[0]
                for row_number, values in enumerate(rows[first_header:], start=first_header + 1):
                    if not any(_has_value(value) for value in values):
                        continue
                    samples.append(
                        WorkbookSampleRow(
                            source_row_number=row_number,
                            values=[
                                _preview_value(value) for value in _trim_trailing_empty(values)
                            ],
                        )
                    )
                    remaining_samples -= 1
                    if remaining_samples == 0:
                        break

            inspected_sheets.append(
                WorkbookSheetInspection(
                    name=worksheet.title,
                    max_row=worksheet.max_row or 0,
                    candidate_header_rows=candidates,
                    header_labels=labels,
                    sample_rows=samples,
                )
            )
    finally:
        workbook.close()

    return WorkbookInspection(sheets=inspected_sheets)


def parse_supplier_workbook(data: bytes, mapping: WorkbookMapping) -> WorkbookParseResult:
    """Parse a supplier workbook using only an explicitly confirmed mapping."""
    workbook, issue = _open_workbook(data)
    if issue is not None:
        return WorkbookParseResult(rows=[], issues=[issue])
    assert workbook is not None

    rows: list[CanonicalQuoteRow] = []
    issues: list[WorkbookRowIssue] = []
    try:
        worksheet = (
            workbook[mapping.worksheet_name]
            if mapping.worksheet_name in workbook.sheetnames
            else None
        )
        if worksheet is None:
            return WorkbookParseResult(
                rows=[],
                issues=[
                    _issue(
                        "worksheet_not_found",
                        "The mapped worksheet is not present in this workbook.",
                        worksheet_name=mapping.worksheet_name,
                    )
                ],
            )

        if worksheet.max_row is not None and worksheet.max_row > MAX_WORKSHEET_ROWS:
            return WorkbookParseResult(
                rows=[],
                issues=[
                    _issue(
                        "row_limit_exceeded",
                        "The worksheet exceeds the supported row limit.",
                        worksheet_name=worksheet.title,
                    )
                ],
            )
        if mapping.header_row > (worksheet.max_row or 0):
            return WorkbookParseResult(
                rows=[],
                issues=[
                    _issue(
                        "header_row_not_found",
                        "The mapped header row is outside the worksheet.",
                        worksheet_name=worksheet.title,
                    )
                ],
            )

        header_values = next(
            worksheet.iter_rows(
                min_row=mapping.header_row,
                max_row=mapping.header_row,
                max_col=MAX_COLUMN_COUNT,
                values_only=True,
            ),
            (),
        )
        header_positions: dict[str, list[int]] = {}
        for index, value in enumerate(header_values, start=1):
            if _has_value(value):
                header_positions.setdefault(_normalize_label(str(value)), []).append(index)

        mapped_columns: dict[WorkbookField, int] = {}
        for field, label in mapping.columns.items():
            positions = header_positions.get(_normalize_label(label), [])
            if len(positions) != 1:
                code = "ambiguous_column" if positions else "column_not_found"
                return WorkbookParseResult(
                    rows=[],
                    issues=[
                        _issue(
                            code,
                            "A mapped column is missing or appears more than once.",
                            worksheet_name=worksheet.title,
                        )
                    ],
                )
            mapped_columns[field] = positions[0]

        merged_ranges = _merged_ranges_for_sheet(data, worksheet.title)
        max_col = max(mapped_columns.values(), default=1)
        for row_number, values in enumerate(
            worksheet.iter_rows(
                min_row=mapping.header_row + 1,
                max_row=worksheet.max_row or mapping.header_row,
                max_col=max_col,
                values_only=True,
            ),
            start=mapping.header_row + 1,
        ):
            row_data = {
                field: values[column - 1] if column <= len(values) else None
                for field, column in mapped_columns.items()
            }
            if not any(_has_value(value) for value in row_data.values()):
                continue
            if _is_repeated_header(row_data, mapping):
                issues.append(
                    _issue(
                        "repeated_header",
                        "A repeated header row was skipped.",
                        worksheet_name=worksheet.title,
                        source_row_number=row_number,
                    )
                )
                continue
            if _row_intersects_merged_mapping(row_number, mapped_columns.values(), merged_ranges):
                issues.append(
                    _issue(
                        "merged_data_cells",
                        "Merged cells overlap mapped data columns; this row needs review.",
                        worksheet_name=worksheet.title,
                        source_row_number=row_number,
                    )
                )
                continue

            parsed, row_issues = _parse_quote_row(
                row_data,
                worksheet_name=worksheet.title,
                row_number=row_number,
                mapping=mapping,
            )
            issues.extend(row_issues)
            if parsed is not None:
                rows.append(parsed)
    finally:
        workbook.close()

    return WorkbookParseResult(rows=rows, issues=issues)


def _open_workbook(data: bytes) -> tuple[Any | None, WorkbookRowIssue | None]:
    if len(data) > MAX_ATTACHMENT_BYTES:
        return None, _issue("file_too_large", "The workbook exceeds the 20 MiB attachment limit.")
    if not data:
        return None, _issue("invalid_workbook", "The attachment is not a readable XLSX workbook.")
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if sum(entry.file_size for entry in entries) > MAX_ZIP_EXPANDED_BYTES:
                return None, _issue(
                    "workbook_expands_too_large", "The workbook expands beyond the safe limit."
                )
            if any(entry.filename.lower().endswith("vbaproject.bin") for entry in entries):
                return None, _issue(
                    "unsupported_macro_workbook", "Macro-enabled workbooks are not supported."
                )
            if archive.testzip() is not None:
                return None, _issue("invalid_workbook", "The XLSX archive contains a corrupt file.")
        workbook = load_workbook(
            BytesIO(data),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (zipfile.BadZipFile, InvalidFileException, OSError, ValueError, KeyError):
        return None, _issue("invalid_workbook", "The attachment is not a readable XLSX workbook.")
    return workbook, None


def _read_preview_rows(worksheet: Worksheet) -> list[tuple[Any, ...]]:
    max_row = min(worksheet.max_row or 0, MAX_HEADER_SCAN_ROWS + MAX_PREVIEW_ROWS)
    if max_row == 0:
        return []
    return [
        tuple(row[:MAX_COLUMN_COUNT])
        for row in worksheet.iter_rows(
            min_row=1,
            max_row=max_row,
            max_col=MAX_COLUMN_COUNT,
            values_only=True,
        )
    ]


def _preview_value(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    return _preview_text(value)


def _preview_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:MAX_CELL_PREVIEW_CHARS]
    return str(value)[:MAX_CELL_PREVIEW_CHARS]


def _has_value(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _trim_trailing_empty(values: tuple[Any, ...]) -> tuple[Any, ...]:
    end = len(values)
    while end and not _has_value(values[end - 1]):
        end -= 1
    return values[:end]


def _normalize_label(label: str) -> str:
    return " ".join(label.strip().casefold().split())


def _merged_ranges_for_sheet(data: bytes, sheet_name: str) -> list[tuple[int, int, int, int]]:
    """Read merge metadata from XLSX XML because read-only worksheets omit it."""
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            rels_root = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            relationships = {
                relationship.attrib["Id"]: relationship.attrib["Target"]
                for relationship in rels_root.findall(f"{{{_PACKAGE_REL_NS}}}Relationship")
            }
            sheet_path: str | None = None
            for sheet in workbook_root.findall(f".//{{{_WORKBOOK_NS}}}sheet"):
                if sheet.attrib.get("name") == sheet_name:
                    target = relationships.get(sheet.attrib.get(f"{{{_REL_NS}}}id", ""))
                    if target is not None:
                        sheet_path = (
                            target.lstrip("/")
                            if target.startswith("/")
                            else posixpath.normpath(posixpath.join("xl", target))
                        )
                    break
            if sheet_path is None:
                return []
            worksheet_root = ElementTree.fromstring(archive.read(sheet_path))
    except (zipfile.BadZipFile, KeyError, ElementTree.ParseError):
        return []

    merged: list[tuple[int, int, int, int]] = []
    for merged_cell in worksheet_root.findall(f".//{{{_WORKBOOK_NS}}}mergeCell"):
        reference = merged_cell.attrib.get("ref")
        if reference is None:
            continue
        try:
            min_col, min_row, max_col, max_row = range_boundaries(reference)
        except ValueError:
            continue
        if (
            min_col is not None
            and min_row is not None
            and max_col is not None
            and max_row is not None
        ):
            merged.append((min_col, min_row, max_col, max_row))
    return merged


def _row_intersects_merged_mapping(
    row_number: int,
    mapped_columns: Any,
    merged_ranges: list[tuple[int, int, int, int]],
) -> bool:
    columns = set(mapped_columns)
    return any(
        min_row <= row_number <= max_row
        and min_col != max_col
        and any(min_col <= column <= max_col for column in columns)
        for min_col, min_row, max_col, max_row in merged_ranges
    )


def _is_repeated_header(row_data: dict[WorkbookField, Any], mapping: WorkbookMapping) -> bool:
    mapped_values = [
        (field, value)
        for field, value in row_data.items()
        if _has_value(value) and field in mapping.columns
    ]
    return bool(mapped_values) and all(
        _normalize_label(str(value)) == _normalize_label(mapping.columns[field])
        for field, value in mapped_values
    )


def _parse_quote_row(
    row_data: dict[WorkbookField, Any],
    *,
    worksheet_name: str,
    row_number: int,
    mapping: WorkbookMapping,
) -> tuple[CanonicalQuoteRow | None, list[WorkbookRowIssue]]:
    issues: list[WorkbookRowIssue] = []
    sku = _optional_text(row_data.get("supplier_sku"))
    product_name = _optional_text(row_data.get("product_name"))
    if sku is None and product_name is None:
        return None, [
            _issue(
                "missing_product_identity",
                "A product name or supplier SKU is required.",
                worksheet_name=worksheet_name,
                source_row_number=row_number,
            )
        ]

    price = _to_decimal(row_data.get("unit_price"))
    if price is None or price <= 0:
        return None, [
            _issue(
                "invalid_price",
                "The unit price is missing or is not a positive decimal value.",
                worksheet_name=worksheet_name,
                source_row_number=row_number,
            )
        ]

    currency = _currency(row_data.get("currency")) or mapping.default_currency
    if currency is None:
        return None, [
            _issue(
                "missing_currency",
                "A currency value or explicit supplier default is required.",
                worksheet_name=worksheet_name,
                source_row_number=row_number,
            )
        ]
    if not re.fullmatch(r"[A-Z]{3}", currency):
        return None, [
            _issue(
                "invalid_currency",
                "The currency must be a three-letter code.",
                worksheet_name=worksheet_name,
                source_row_number=row_number,
            )
        ]

    purchase_unit = _optional_text(row_data.get("unit")) or mapping.default_purchase_unit
    if purchase_unit is None:
        return None, [
            _issue(
                "missing_unit",
                "A purchase unit or explicit supplier default is required.",
                worksheet_name=worksheet_name,
                source_row_number=row_number,
            )
        ]
    if len(purchase_unit) > 64:
        return None, [
            _issue(
                "invalid_unit",
                "The purchase unit exceeds the supported length.",
                worksheet_name=worksheet_name,
                source_row_number=row_number,
            )
        ]

    quantity_tiers = _parse_quantity_tiers(row_data.get("quantity_tiers"))
    if quantity_tiers is None:
        issues.append(
            _issue(
                "invalid_quantity_tiers",
                "Quantity tiers could not be parsed and need review.",
                worksheet_name=worksheet_name,
                source_row_number=row_number,
            )
        )
        quantity_tiers = []

    moq: Decimal | None = None
    raw_moq = row_data.get("minimum_order_quantity")
    if _has_value(raw_moq):
        moq = _to_decimal(raw_moq)
        if moq is None or moq < 0:
            issues.append(
                _issue(
                    "invalid_moq",
                    "Minimum order quantity is not a non-negative decimal value.",
                    worksheet_name=worksheet_name,
                    source_row_number=row_number,
                )
            )
            moq = None

    valid_until = _parse_date(row_data.get("valid_until"))
    if _has_value(row_data.get("valid_until")) and valid_until is None:
        issues.append(
            _issue(
                "invalid_valid_until",
                "Quote validity date could not be parsed.",
                worksheet_name=worksheet_name,
                source_row_number=row_number,
            )
        )

    try:
        parsed = CanonicalQuoteRow(
            worksheet_name=worksheet_name,
            source_row_number=row_number,
            supplier_sku=sku,
            product_name=product_name,
            specification=_optional_text(row_data.get("specification")),
            unit_price=price,
            currency=currency,
            purchase_unit=purchase_unit,
            minimum_order_quantity=moq,
            quantity_tiers=quantity_tiers,
            lead_time=_optional_text(row_data.get("lead_time")),
            valid_until=valid_until,
            extraction_confidence=1.0,
        )
    except ValueError:
        return None, [
            _issue(
                "invalid_quote_row",
                "The row contains values outside the supported quote format.",
                worksheet_name=worksheet_name,
                source_row_number=row_number,
            )
        ]
    return parsed, issues


def _parse_quantity_tiers(value: Any) -> list[QuantityTier] | None:
    if not _has_value(value):
        return []
    if not isinstance(value, str):
        return None
    tiers: list[QuantityTier] = []
    for item in re.split(r"[;\n]+", value):
        if not item.strip():
            continue
        match = _QUANTITY_TIER_RE.fullmatch(item)
        if match is None:
            return None
        try:
            tiers.append(
                QuantityTier(
                    minimum_quantity=Decimal(match.group(1)),
                    unit_price=Decimal(match.group(2)),
                )
            )
        except (InvalidOperation, ValueError):
            return None
    return tiers


def _to_decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        result = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return None
    return result if result.is_finite() else None


def _currency(value: Any) -> str | None:
    text = _optional_text(value)
    return text.upper() if text is not None else None


def _optional_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str):
        return None
    try:
        return date.fromisoformat(value.strip())
    except ValueError:
        return None


def _issue(
    code: str,
    message: str,
    *,
    worksheet_name: str | None = None,
    source_row_number: int | None = None,
) -> WorkbookRowIssue:
    return WorkbookRowIssue(
        worksheet_name=worksheet_name,
        source_row_number=source_row_number,
        code=code,
        message=message[:160],
    )
