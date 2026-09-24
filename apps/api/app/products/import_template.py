"""The downloadable sample catalogue behind the Products page's "Download
sample" control, in each format `app/products/importer.py` reads.

Built from the importer's own `_KNOWN_FIELDS` rather than kept as static
files, so the sample's columns cannot drift from the ones the parser
actually reads -- `tests/unit/test_product_importer.py` feeds every format
back through `parse_import` and requires zero row errors.
"""

import csv
import io
import json
from dataclasses import dataclass
from typing import Any, Literal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.products.importer import _KNOWN_FIELDS, XLSX_MIME_TYPE

TemplateFormat = Literal["csv", "xlsx", "json"]

_REQUIRED = frozenset({"external_id", "name"})

# One row per shape worth showing: a fully filled product, a minimal one
# (slug left out to be generated from `name`), and a non-default
# availability. `attributes`/`metadata` are native objects here; the CSV and
# XLSX writers encode them as the JSON-in-a-cell the parser expects.
_SAMPLE_ROWS: tuple[dict[str, Any], ...] = (
    {
        "external_id": "SKU-001",
        "name": "Classic White T-Shirt",
        "slug": "classic-white-t-shirt",
        "description": "100% cotton crew-neck tee.",
        "category": "Apparel",
        "price": 199000,
        "currency": "VND",
        "attributes": {"size": "M", "color": "white"},
        "availability": "in_stock",
        "stock_quantity": 50,
        "image_url": "https://example.com/images/sku-001.jpg",
        "product_url": "https://example.com/products/sku-001",
        "is_active": True,
        "metadata": {"supplier": "ACME"},
    },
    {
        "external_id": "SKU-002",
        "name": "Slim Fit Jeans",
        "category": "Apparel",
        "price": 450000,
        "currency": "VND",
    },
    {
        "external_id": "SKU-003",
        "name": "Leather Belt",
        "category": "Accessories",
        "price": 250000.5,
        "currency": "VND",
        "availability": "out_of_stock",
        "stock_quantity": 0,
    },
)

# The XLSX's second sheet -- the importer only reads the first, so this is
# documentation that travels with the file, not data.
_COLUMN_NOTES: dict[str, str] = {
    "external_id": "Your SKU or product code. Unique; importing it again updates the product.",
    "name": "Product name.",
    "slug": "Optional. Generated from name when blank.",
    "description": "Free text.",
    "category": "Free text.",
    "price": "Number, at most 2 decimal places.",
    "currency": "3-letter code, e.g. VND or USD.",
    "attributes": 'JSON object, e.g. {"size": "M"}.',
    "availability": "in_stock (default), out_of_stock, preorder or discontinued.",
    "stock_quantity": "Whole number.",
    "image_url": "Link to an image.",
    "product_url": "Link to the product page.",
    "is_active": "true (default) or false.",
    "metadata": "JSON object for your own data.",
}


@dataclass(frozen=True, slots=True)
class ImportTemplate:
    content: bytes
    media_type: str
    filename: str


def _cell_text(field: str, value: Any) -> Any:
    """A value as a flat cell: JSON fields encoded, absent fields blank."""
    if value is None:
        return None
    if field in ("attributes", "metadata"):
        return json.dumps(value, ensure_ascii=False)
    return value


def _csv_cell(field: str, value: Any) -> Any:
    cell = _cell_text(field, value)
    if cell is None:
        return ""
    if isinstance(cell, bool):
        return "true" if cell else "false"
    return cell


def _csv_bytes() -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(_KNOWN_FIELDS)
    for row in _SAMPLE_ROWS:
        writer.writerow([_csv_cell(name, row.get(name)) for name in _KNOWN_FIELDS])
    # BOM so Excel opens the file as UTF-8 rather than the system codepage;
    # the importer reads with `utf-8-sig` and strips it again.
    return buffer.getvalue().encode("utf-8-sig")


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Products"
    sheet.append(list(_KNOWN_FIELDS))
    for row in _SAMPLE_ROWS:
        sheet.append([_cell_text(name, row.get(name)) for name in _KNOWN_FIELDS])

    required_fill = PatternFill("solid", fgColor="FFF2CC")
    for column, name in enumerate(_KNOWN_FIELDS, start=1):
        cell = sheet.cell(row=1, column=column)
        cell.font = Font(bold=True)
        if name in _REQUIRED:
            cell.fill = required_fill
        sheet.column_dimensions[get_column_letter(column)].width = max(14, len(name) + 4)
    sheet.freeze_panes = "A2"

    notes = workbook.create_sheet("Columns")
    notes.append(["Column", "Required", "Format"])
    for cell in notes[1]:
        cell.font = Font(bold=True)
    for name in _KNOWN_FIELDS:
        notes.append([name, "yes" if name in _REQUIRED else "no", _COLUMN_NOTES[name]])
    notes.column_dimensions["A"].width = 16
    notes.column_dimensions["C"].width = 80

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _json_bytes() -> bytes:
    return json.dumps(list(_SAMPLE_ROWS), ensure_ascii=False, indent=2).encode("utf-8")


def build_import_template(fmt: TemplateFormat) -> ImportTemplate:
    filename = f"product-import-sample.{fmt}"
    if fmt == "csv":
        return ImportTemplate(_csv_bytes(), "text/csv", filename)
    if fmt == "xlsx":
        return ImportTemplate(_xlsx_bytes(), XLSX_MIME_TYPE, filename)
    return ImportTemplate(_json_bytes(), "application/json", filename)
