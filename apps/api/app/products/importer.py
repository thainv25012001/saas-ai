"""Task 3: turning an uploaded CSV, XLSX or JSON catalogue into product rows and
vectors -- `docs/PHASE-5.md` §5.

Two things this module deliberately does NOT do, because Tasks 1 and 2
already built them:

- It never builds its own idempotence. `ProductService.upsert_many`'s
  `UNIQUE (organization_id, external_id)` / `ON CONFLICT` is what makes a
  re-import of the same file a no-op rather than a duplicate; this module
  only has to get the right `ProductInput` rows to it.
- It never embeds inline. `app/products/embedding.py::embed_and_store` is
  the second phase; every row this module upserts is handed to
  `needs_reembedding` afterward, and only the ones that actually need a
  vector (never embedded, or embedded from text that no longer matches --
  see that module's docstring for the three states) go to `embed_and_store`.
  A re-import whose content did not change re-embeds nothing.

**Two phases in one arq job: product data first, embedding afterwards
(final review I3).** Phase 1 upserts every chunk, each committed on its
own, with no embedding provider involved. Phase 2 then embeds the rows
phase 1 reported as needing it, in separate transactions. The split is
what keeps an embedding outage from costing the customer their prices and
stock: `docs/PHASE-5.md` §5/§8 treat "metadata written, embedding not yet
recomputed" as a normal state, search serves unembedded rows through
filters and full text, and the dashboard badges them "Not yet searchable
by meaning". So when the provider fails, phase 2 stops (it does not keep
calling a provider that is down), leaves the remaining rows unembedded,
and records a warning on the import -- the import still completes with
every valid row in the table. Re-importing the file retries the embedding,
because an unembedded row still answers `needs_reembedding`. One job, not
two queued jobs: a second job would add an ordering dependency (it must
not start before the upserts commit) for nothing phase 2 cannot already
do here.

**Per-row validation, chunked writes, no giant single transaction and no
one-row-per-transaction either.** Every row is validated independently in
Python (`_build_product_input`, via `ProductInput`) before anything touches
the database, so one bad price never prevents the other 3,999 rows from
being considered at all -- this is what makes "batch-level success"
possible in the first place, not a database-level retry of anything.
Valid rows are then written in chunks of `_IMPORT_CHUNK_SIZE`, each chunk
its own `ProductService.upsert_many` call (one INSERT ... ON CONFLICT
statement) in its own transaction. A transaction per row would turn a
4,000-row import into 4,000+ round trips for no isolation benefit --
`upsert_many` already validated every row in Python, so the only thing left
that can fail per-chunk is a genuine database error (a connection blip, a
constraint this module did not anticipate), and chunking bounds both how
much work such a failure discards (at most `_IMPORT_CHUNK_SIZE` rows, not
the whole file) and how large any single SQL statement gets (bounded
parameter count, bounded lock duration).
"""

import asyncio
import csv
import io
import json
import re
import uuid
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import ValidationError as PydanticValidationError
from sqlalchemy import select
from sqlalchemy.exc import DBAPIError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.errors import AppError, NotFoundError, format_validation_errors
from app.core.ids import uuid7
from app.core.logging import get_logger
from app.core.tenancy import TenantContext, tenant_session
from app.db.models import Product, ProductImport, ProductImportStatus
from app.products.embedding import (
    embed_and_store,
    needs_reembedding,
    needs_reembedding_clause,
)
from app.products.schemas import ProductInput
from app.products.service import ProductService
from app.rag.ingest import bounded_error_message

logger = get_logger(__name__)


class UnsupportedImportType(AppError):
    code = "unsupported_import_type"
    status_code = 422


class ImportParseError(AppError):
    """A whole-file structural problem, not a bad row -- a CSV missing a
    required column, or a JSON document whose top level is not an array.
    Raised before any per-row work starts, and deliberately not caught
    per-row: there is no row to attach the error to, and reporting the same
    "missing column" failure once per row would bury the one thing the
    customer actually needs to fix under thousands of copies of it."""

    code = "invalid_import_file"
    status_code = 422


# ---------------------------------------------------------------------------
# Mime type resolution -- the same "browser says nothing useful, fall back
# to the extension" shape as app/rag/extract.py::resolve_mime_type, kept as
# its own small copy rather than generalising that function: the two
# problems only look alike from a distance (a completely different
# EXTENSION_MIME_TYPES table, a completely different SUPPORTED set), and
# threading a table through a shared helper would couple an unrelated
# document-upload module to this one for a dozen lines saved.
# ---------------------------------------------------------------------------

XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SUPPORTED_IMPORT_MIME_TYPES: frozenset[str] = frozenset(
    {"text/csv", "application/json", XLSX_MIME_TYPE}
)

_GENERIC_MIME_TYPES: frozenset[str] = frozenset(
    {"", "application/octet-stream", "binary/octet-stream"}
)

_EXTENSION_MIME_TYPES: dict[str, str] = {
    ".csv": "text/csv",
    ".json": "application/json",
    ".xlsx": XLSX_MIME_TYPE,
}

# Types browsers and operating systems report for a file that is really a
# `.csv` (final review I4). Chromium and Firefox on Windows take `.csv`'s
# type from the registry, which is `application/vnd.ms-excel` whenever Excel
# is installed -- the most common source of a customer's catalogue -- and
# other platforms report `text/plain` or one of the older CSV spellings.
# For these, and only when the filename ends in `.csv`, the extension is
# authoritative. Anything else reported is still believed over the
# extension, as before.
_CSV_ALIAS_MIME_TYPES: frozenset[str] = frozenset(
    {
        "application/vnd.ms-excel",
        "text/plain",
        "text/x-csv",
        "application/csv",
        "application/x-csv",
        "text/comma-separated-values",
    }
)


def resolve_import_mime_type(reported: str | None, filename: str | None) -> str:
    """What to treat this upload as -- see `resolve_mime_type`'s docstring
    in app/rag/extract.py for why the reported type wins whenever it is
    anything other than empty/octet-stream, and the extension is only ever
    a fallback for the browsers that report nothing useful. The one
    exception is `_CSV_ALIAS_MIME_TYPES` on a `.csv` file. MIME parameters
    (`text/csv; charset=utf-8`) are stripped before any comparison."""
    reported = reported or ""
    base = reported.split(";", 1)[0].strip().lower()
    suffix_type: str | None = None
    if filename:
        _, _, suffix = filename.rpartition(".")
        suffix_type = _EXTENSION_MIME_TYPES.get(f".{suffix.lower()}")
    if base in _GENERIC_MIME_TYPES:
        return suffix_type or reported
    if base in _CSV_ALIAS_MIME_TYPES and suffix_type == "text/csv":
        return suffix_type
    return base


# ---------------------------------------------------------------------------
# Row parsing
# ---------------------------------------------------------------------------

# Every ProductInput field an import row may supply. Deliberately excludes
# `embedding`/`embedding_model`: an import's first phase always supplies
# neither (see the module docstring), so even a JSON source file that names
# them is not read -- `_build_product_input` below never looks at those two
# keys in `raw`, whatever it contains.
_KNOWN_FIELDS: tuple[str, ...] = (
    "external_id",
    "name",
    "slug",
    "description",
    "category",
    "price",
    "currency",
    "attributes",
    "availability",
    "stock_quantity",
    "image_url",
    "product_url",
    "is_active",
    "metadata",
)

# `attributes`/`metadata` are jsonb: arbitrary shape, so a flat CSV cell
# carries them as a JSON-encoded string (e.g. `{"seats": 5, "fuel": "hybrid"}`)
# rather than this module inventing a column-splatting convention (a
# `attributes.seats` header naming scheme, say) that Task 1 never specified
# and nothing downstream expects. A JSON import supplies these as native
# objects already, so `_build_product_input` accepts either shape uniformly
# -- see its docstring.
_JSON_FIELDS: frozenset[str] = frozenset({"attributes", "metadata"})

# Optional scalar fields where an empty CSV cell means "not supplied" (None)
# rather than the literal empty string -- required for `price`/`stock_quantity`
# in particular, since pydantic would otherwise try (and fail) to parse ""
# as a Decimal/int. `external_id`/`name`/`slug` are deliberately NOT in this
# set: those are required, and an empty cell for one of them must surface as
# "external_id: at least 1 character", not silently become `None` and fail
# with a less legible "not a valid string" instead.
_BLANK_MEANS_NONE_FIELDS: frozenset[str] = frozenset(
    {"description", "category", "price", "currency", "stock_quantity", "image_url", "product_url"}
)

# `availability`/`is_active` both carry defaults on `ProductInput`
# (IN_STOCK / True); an empty cell should take that default rather than be
# forced through pydantic as an empty string or empty-string-turned-None,
# either of which would fail validation instead of falling back.
_DEFAULTED_FIELDS: frozenset[str] = frozenset({"availability", "is_active"})


@dataclass(frozen=True, slots=True)
class RowError:
    """One row's reason for not being imported. `row` is 1-based and, for a
    CSV, header-aware: row 1 is the header, so the first data row is row 2
    -- the number a human looking at their own spreadsheet would point at,
    not a 0-based data index that is off by one from what they see."""

    row: int
    external_id: str | None
    message: str

    def to_dict(self) -> dict[str, Any]:
        return {"row": self.row, "external_id": self.external_id, "message": self.message}


@dataclass(frozen=True, slots=True)
class ParsedRow:
    row: int
    product: ProductInput


@dataclass(slots=True)
class ParseResult:
    rows: list[ParsedRow] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    total_rows: int = 0


def _slugify(name: str) -> str:
    """A `slug` a customer's own catalogue almost never carries -- it is an
    internal concept, not a field a source ERP or spreadsheet has any
    reason to export. Auto-generated from `name` when the row omits it (or
    leaves it blank) so a slug-less import is not simply a 4,000-row hard
    error; a row that does supply one is trusted as-is (`ProductInput` puts
    no format constraint on it beyond length)."""
    slug = re.sub(r"[^a-z0-9]+", "-", name.strip().lower()).strip("-")
    return slug or "product"


def _external_id_of(raw: dict[str, Any]) -> str | None:
    value = raw.get("external_id")
    if value is None or value == "":
        return None
    return str(value).strip()


def _describe_row_error(exc: Exception) -> str:
    if isinstance(exc, PydanticValidationError):
        return format_validation_errors(exc.errors())
    if isinstance(exc, json.JSONDecodeError):
        return f"invalid JSON: {exc}"
    return str(exc)


def _row_error(row: int, raw: dict[str, Any], exc: Exception) -> RowError:
    return RowError(row=row, external_id=_external_id_of(raw), message=_describe_row_error(exc))


def _build_product_input(raw: dict[str, Any]) -> ProductInput:
    """Build one row's `ProductInput` from its raw field values.

    `raw`'s values may be strings (every CSV cell) or already-native JSON
    types (a JSON import's own objects/numbers/booleans) -- both are
    accepted uniformly here rather than this module maintaining two
    separate builders, one per source format. Raises pydantic's
    `ValidationError` or `json.JSONDecodeError`; both are caught by this
    module's callers and turned into a `RowError`, never left to escape and
    abort the rest of the file.
    """
    fields: dict[str, Any] = {}
    for name in _KNOWN_FIELDS:
        if name not in raw:
            continue
        value = raw[name]
        if name in _JSON_FIELDS:
            if value in (None, ""):
                continue
            fields[name] = json.loads(value) if isinstance(value, str) else value
        elif name in _BLANK_MEANS_NONE_FIELDS:
            fields[name] = None if value in (None, "") else value
        elif name in _DEFAULTED_FIELDS:
            if value not in (None, ""):
                fields[name] = value
        else:
            # external_id/name/slug: pass through even when blank, so a
            # missing required field fails with pydantic's own "at least 1
            # character" message rather than being silently turned into
            # `None` first (see `_BLANK_MEANS_NONE_FIELDS`'s docstring).
            fields[name] = "" if value is None else value

    if not fields.get("slug"):
        fields["slug"] = _slugify(str(fields.get("name") or ""))

    return ProductInput(**fields)


def _parse_table(
    header: Iterable[str | None], rows: Iterable[tuple[int, dict[str | None, Any]]], kind: str
) -> ParseResult:
    """The row loop CSV and XLSX share: one required-column check for the
    whole file, then every row validated on its own. `rows` carries each
    row's own number -- the one the customer sees in their file -- because
    only the source knows it (an XLSX skips blank rows without renumbering
    the ones after them)."""
    missing = {"external_id", "name"} - set(header)
    if missing:
        raise ImportParseError(
            f"{kind} is missing required column(s): {', '.join(sorted(missing))}"
        )

    result = ParseResult()
    for row_number, raw_row in rows:
        result.total_rows += 1
        # A short row (fewer cells than the header) fills the missing keys
        # with `None` (csv.DictReader's `restval`); a long one adds a
        # `None` key holding the overflow. Neither is a value any field
        # above expects, so both are normalised the same way every other
        # blank cell is.
        raw = {
            key: ("" if value is None else value)
            for key, value in raw_row.items()
            if key is not None
        }
        try:
            product = _build_product_input(raw)
        except (PydanticValidationError, json.JSONDecodeError) as exc:
            result.errors.append(_row_error(row_number, raw, exc))
            continue
        result.rows.append(ParsedRow(row=row_number, product=product))
    return result


def _parse_csv(data: bytes) -> ParseResult:
    # `utf-8-sig`: Excel's "CSV UTF-8" export prepends a BOM, which -- left
    # in place -- would land inside the *value* of the first header cell
    # ("﻿external_id"), silently making that column invisible to every
    # lookup below and turning "CSV is missing required column(s):
    # external_id" into a support ticket about a column that is right there
    # on screen.
    reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig")))
    return _parse_table(
        reader.fieldnames or [],
        enumerate(reader, start=2),  # the header consumes row 1
        "CSV",
    )


def _xlsx_cell_text(value: Any) -> str:
    """A cell as the text the same value would be in a CSV export, so an
    XLSX row goes through exactly the rules a CSV row does. Excel keeps
    numbers as numbers: a SKU typed as 1001 arrives as an int (which the
    string-typed `external_id` would reject) and a quantity of 50 often as
    50.0 (which `stock_quantity` would reject as "not a whole number" only
    by luck of float formatting). Booleans become the `true`/`false` a CSV
    would carry."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _parse_xlsx(data: bytes) -> ParseResult:
    """The first sheet, row 1 the header. Read-only (streamed, not the
    whole workbook model in memory) and `data_only`, so a formula cell
    imports the value Excel last computed for it rather than its formula
    text. Rows with no value in any cell are skipped rather than reported:
    Excel routinely leaves formatted-but-empty rows below the data, and
    listing each as a failed row would bury the real ones."""
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            cells = [
                [_xlsx_cell_text(value) for value in row]
                for row in workbook.worksheets[0].iter_rows(values_only=True)
            ]
        finally:
            workbook.close()
    except Exception as exc:
        # Anything openpyxl raises on these bytes -- not a zip, a zip that
        # is not a workbook, a workbook with no sheet -- means the same
        # thing to the customer, and none of it is one row's fault.
        raise ImportParseError(
            "file is not a readable .xlsx workbook (an older .xls must be saved as .xlsx)"
        ) from exc

    if not cells:
        raise ImportParseError("XLSX has no header row")
    header: list[str | None] = [name or None for name in cells[0]]
    rows = (
        (row_number, dict(zip(header, row, strict=False)))
        for row_number, row in enumerate(cells[1:], start=2)
        if any(row)
    )
    return _parse_table(header, rows, "XLSX")


def _parse_json(data: bytes) -> ParseResult:
    try:
        payload = json.loads(data)
    except json.JSONDecodeError as exc:
        raise ImportParseError(f"invalid JSON: {exc}") from exc
    if not isinstance(payload, list):
        raise ImportParseError("JSON import must be an array of product objects")

    result = ParseResult()
    for index, item in enumerate(payload):
        row_number = index + 1  # no header row to offset past, unlike CSV
        result.total_rows += 1
        if not isinstance(item, dict):
            result.errors.append(
                RowError(
                    row=row_number, external_id=None, message="each item must be a JSON object"
                )
            )
            continue
        try:
            product = _build_product_input(item)
        except (PydanticValidationError, json.JSONDecodeError) as exc:
            result.errors.append(_row_error(row_number, item, exc))
            continue
        result.rows.append(ParsedRow(row=row_number, product=product))
    return result


def parse_import(data: bytes, mime_type: str) -> ParseResult:
    """Parse an uploaded catalogue into valid rows plus per-row errors.
    Never raises for a bad *row* -- only for a whole-file structural
    problem (`ImportParseError`) or a mime type this module was never
    taught (`UnsupportedImportType`, which the API layer should already
    have rejected via `SUPPORTED_IMPORT_MIME_TYPES` before this is ever
    called; kept here too so a caller that skips that check still fails
    loudly instead of silently returning zero rows)."""
    if mime_type == "text/csv":
        return _parse_csv(data)
    if mime_type == "application/json":
        return _parse_json(data)
    if mime_type == XLSX_MIME_TYPE:
        return _parse_xlsx(data)
    raise UnsupportedImportType(f"unsupported import type '{mime_type}'")


def _dedupe_by_external_id(rows: list[ParsedRow]) -> tuple[list[ParsedRow], list[RowError]]:
    """Keep the LAST occurrence of each `external_id` in file order --
    ordinary upsert semantics: a later row for the same SKU wins, the same
    way a second import of the same SKU would.

    This has to happen before chunking, not per chunk: `ProductService.
    upsert_many` issues one `INSERT ... ON CONFLICT` statement per call, and
    Postgres refuses to let a single statement affect the same conflict
    target twice ("ON CONFLICT DO UPDATE command cannot affect row a second
    time") -- two rows sharing an `external_id` inside the same chunk would
    abort that whole chunk, not just one of the two rows. Deduping globally,
    before the file is ever split into chunks, also means which rows get
    reported as a duplicate never depends on `_IMPORT_CHUNK_SIZE` -- an
    implementation detail the customer's file should not have to know
    about to predict its own import result.
    """
    by_external_id: dict[str, ParsedRow] = {}
    duplicate_errors: list[RowError] = []
    for parsed in rows:
        external_id = parsed.product.external_id
        prior = by_external_id.get(external_id)
        if prior is not None:
            duplicate_errors.append(
                RowError(
                    row=prior.row,
                    external_id=external_id,
                    message=(
                        f"external_id '{external_id}' also appears on row {parsed.row}; "
                        "only the later occurrence was imported"
                    ),
                )
            )
        by_external_id[external_id] = parsed
    return list(by_external_id.values()), duplicate_errors


# Rows per `upsert_many` call / transaction. Large enough that a multi
# thousand row catalogue lands in a handful of round trips, not one per
# row; small enough that one chunk's parameter count stays far under
# Postgres's 65535-per-statement limit (this table's ~15 upsert columns *
# 500 rows is ~7500) and that one failing chunk discards a bounded slice of
# the file rather than the whole thing.
_IMPORT_CHUNK_SIZE = 500

# Ceiling on `ProductImportService.list_imports`'s `limit` -- same
# reasoning as `ProductService`'s own `_MAX_LIST_LIMIT`. Each record can
# carry a per-row error list thousands long, so an unbounded page is a
# large serialize, not just a large scan.
_MAX_LIST_LIMIT = 100


def _chunks[T](rows: list[T], size: int) -> list[list[T]]:
    return [rows[index : index + size] for index in range(0, len(rows), size)]


async def _upsert_chunk(tenant: TenantContext, chunk: list[ParsedRow]) -> list[uuid.UUID]:
    """Phase 1 for one chunk: upsert it in its own committed transaction and
    return the ids of the rows that need a vector.

    `needs_reembedding` (app/products/embedding.py) is what keeps a
    re-import cheap: a row whose upsert changed nothing embeddable (a
    price-only sync riding along in the same file, or a genuinely identical
    re-import) comes back with its previous, still-current embedding and is
    left out here, so phase 2 never embeds it. A row that is still
    unembedded from an earlier import whose embedding failed comes back
    with no hash and IS returned -- that is how re-importing retries it.
    """
    async with tenant_session(tenant) as session:
        results = await ProductService(session, tenant).upsert_many(
            [parsed.product for parsed in chunk]
        )
        return [product.id for product in results if needs_reembedding(product)]


async def _upsert_rows_individually(
    tenant: TenantContext, chunk: list[ParsedRow]
) -> tuple[list[uuid.UUID], int, list[RowError]]:
    """Fallback for a chunk whose batched upsert failed at the database for
    a reason `ProductInput` did not catch -- re-run one row at a time, each
    its own transaction, so the row(s) actually at fault are the only ones
    reported failed.

    This is what makes "one bad row does not abort the batch" hold for
    *any* database-level failure, not just the ones Python-level validation
    happens to anticipate today: convicting all 500 rows in a chunk for one
    row's fault would be exactly the failure mode the brief exists to
    prevent, just moved one layer down from where per-row validation closes
    it. Deliberately only reached from the chunk-level `except DBAPIError`
    in `run_import` -- a clean import never pays a per-row round trip.

    Catches `DBAPIError` only. No embedding provider is called in phase 1
    any more, and anything else escaping an upsert is not one row's fault,
    so it propagates and fails the import rather than being retried 500
    times (final review I3: the old catch-all turned one provider outage
    into a single-row provider call per row).

    If every row in the chunk fails identically (a connection outage, not
    one row's fault) this still produces the correct outcome: every row is
    reported failed, each with its own accurate message.
    """
    pending: list[uuid.UUID] = []
    succeeded = 0
    errors: list[RowError] = []
    for parsed in chunk:
        try:
            pending.extend(await _upsert_chunk(tenant, [parsed]))
            succeeded += 1
        except DBAPIError as exc:
            errors.append(
                RowError(
                    row=parsed.row,
                    external_id=parsed.product.external_id,
                    message=bounded_error_message(exc),
                )
            )
    return pending, succeeded, errors


async def _embed_rows(tenant: TenantContext, product_ids: list[uuid.UUID]) -> None:
    """Phase 2 for one batch: embed the rows phase 1 upserted, in their own
    transaction.

    Re-selected rather than reusing phase 1's ORM objects (whose session is
    gone), with both tenancy layers: the explicit `organization_id`
    predicate here, and RLS on the session. `needs_reembedding_clause`
    re-checks each row, so a row something else embedded in between is
    skipped. `FOR UPDATE` holds the rows until the vectors are written, so
    an upsert that changes a row's text concurrently waits and then marks
    the fresh vector stale against its new text, rather than this write
    certifying old text as current after that upsert has landed.
    """
    async with tenant_session(tenant) as session:
        result = await session.execute(
            select(Product)
            .where(
                Product.organization_id == tenant.organization_id,
                Product.id.in_(product_ids),
                needs_reembedding_clause(),
            )
            .order_by(Product.id)
            .with_for_update()
        )
        await embed_and_store(session, tenant, list(result.scalars().all()))


async def _embed_imported_rows(
    tenant: TenantContext, product_import_id: uuid.UUID, product_ids: list[uuid.UUID]
) -> str | None:
    """Phase 2: embed `product_ids` in batches of `_IMPORT_CHUNK_SIZE`.
    Returns `None` when every batch was embedded, or a warning for the
    import record when one failed.

    Stops at the first failed batch: `embed_and_store` already retried it
    (`embed_batched`), so the provider is down or throttling, and calling
    it again for every remaining batch would only make that worse and eat
    into the job's timeout. The rows already committed in phase 1 stay as
    they are -- present, and simply not yet searchable by meaning.
    """
    for index, batch in enumerate(_chunks(product_ids, _IMPORT_CHUNK_SIZE)):
        try:
            await _embed_rows(tenant, batch)
        except Exception as exc:
            unembedded = len(product_ids) - index * _IMPORT_CHUNK_SIZE
            logger.warning(
                "product_import_embedding_failed",
                product_import_id=str(product_import_id),
                organization_id=str(tenant.organization_id),
                unembedded_rows=unembedded,
                error=bounded_error_message(exc),
            )
            return (
                f"All valid rows were imported, but {unembedded} could not be embedded "
                f"for search by meaning ({bounded_error_message(exc)}). They can still be "
                "found by name, category and filters. Importing the file again retries "
                "the embedding."
            )
    return None


class ProductImportService:
    """CRUD for `product_imports` -- `DocumentService`'s status-transition
    idiom, applied to a row that also carries counts and a per-row error
    list `Document` has no equivalent of."""

    def __init__(self, session: AsyncSession, tenant: TenantContext) -> None:
        self.session = session
        self.tenant = tenant

    async def create(self, *, filename: str | None, mime_type: str) -> ProductImport:
        record = ProductImport(
            id=uuid7(),
            organization_id=self.tenant.organization_id,
            filename=filename,
            mime_type=mime_type,
            status=ProductImportStatus.PENDING,
        )
        self.session.add(record)
        await self.session.flush()
        return record

    async def get(self, product_import_id: uuid.UUID) -> ProductImport:
        result = await self.session.execute(
            select(ProductImport).where(
                ProductImport.id == product_import_id,
                ProductImport.organization_id == self.tenant.organization_id,
            )
        )
        record = result.scalar_one_or_none()
        if record is None:
            # Cross-tenant lookups fail the same way a nonexistent id does
            # -- see the identical note on DocumentService.get.
            raise NotFoundError("product import not found")
        return record

    async def list_imports(self, *, limit: int = 20, offset: int = 0) -> list[ProductImport]:
        """This organization's imports, newest first -- the dashboard's
        import history. Clamped like `ProductService.list_products`; `id`
        breaks `created_at` ties for a stable page order."""
        limit = max(1, min(limit, _MAX_LIST_LIMIT))
        offset = max(0, offset)
        result = await self.session.execute(
            select(ProductImport)
            .where(ProductImport.organization_id == self.tenant.organization_id)
            .order_by(ProductImport.created_at.desc(), ProductImport.id.desc())
            .limit(limit)
            .offset(offset)
        )
        return list(result.scalars().all())

    async def mark_processing(self, product_import_id: uuid.UUID) -> ProductImport:
        record = await self.get(product_import_id)
        record.status = ProductImportStatus.PROCESSING
        await self.session.flush()
        return record

    async def mark_completed(
        self,
        product_import_id: uuid.UUID,
        *,
        total_rows: int,
        succeeded_count: int,
        failed_count: int,
        errors: list[dict[str, Any]],
        warning: str | None = None,
    ) -> ProductImport:
        """`warning` lands in `error`: on a completed import it is a
        non-fatal problem with the import as a whole (today, only phase 2's
        embedding failure), not the whole-file failure it means on a
        failed one."""
        record = await self.get(product_import_id)
        record.status = ProductImportStatus.COMPLETED
        record.total_rows = total_rows
        record.succeeded_count = succeeded_count
        record.failed_count = failed_count
        record.errors = errors
        record.error = warning
        record.completed_at = datetime.now(UTC)
        await self.session.flush()
        return record

    async def mark_failed(self, product_import_id: uuid.UUID, message: str) -> ProductImport:
        record = await self.get(product_import_id)
        record.status = ProductImportStatus.FAILED
        record.error = message
        record.completed_at = datetime.now(UTC)
        await self.session.flush()
        return record


# ---------------------------------------------------------------------------
# Uploaded-file storage -- app/rag/storage.py's shape, under its own
# subdirectory rather than reusing that module's (document-keyed) functions
# directly: an import id and a document id are both plain uuids, so nothing
# would break by accident today, but nothing here should depend on the two
# id spaces never colliding either.
# ---------------------------------------------------------------------------


def _import_bytes_path(organization_id: uuid.UUID, product_import_id: uuid.UUID) -> Path:
    return (
        Path(get_settings().upload_dir)
        / "product_imports"
        / str(organization_id)
        / str(product_import_id)
    )


def store_import_bytes(
    organization_id: uuid.UUID, product_import_id: uuid.UUID, data: bytes
) -> Path:
    path = _import_bytes_path(organization_id, product_import_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return path


def load_import_bytes(organization_id: uuid.UUID, product_import_id: uuid.UUID) -> bytes:
    return _import_bytes_path(organization_id, product_import_id).read_bytes()


async def enqueue_product_import(product_import_id: uuid.UUID, organization_id: uuid.UUID) -> None:
    """Queue `import_products_task` -- `app/rag/queue.py::enqueue_ingest`'s
    counterpart for Task 3, kept beside the rest of this module rather than
    a new `app/products/queue.py` for one two-line function.

    Imported lazily, not at module level: `app.workers.tasks` imports
    `ProductImportService`/`run_import`/`load_import_bytes` from *this*
    module to define `import_products_task` in the first place, so a
    top-level `from app.workers.tasks import import_products_task` here
    would be a real import cycle, not just an unused one.
    """
    from app.workers.enqueue import enqueue
    from app.workers.tasks import import_products_task

    await enqueue(
        import_products_task,
        organization_id=str(organization_id),
        product_import_id=str(product_import_id),
    )


# ---------------------------------------------------------------------------
# Orchestration -- the arq job body. `app/workers/tasks.py::import_products_task`
# is the thin arq entry point around this, exactly the way `ingest_document_task`
# wraps `app/rag/ingest.py::ingest_document`.
# ---------------------------------------------------------------------------


async def run_import(
    tenant: TenantContext, product_import_id: uuid.UUID, data: bytes, mime_type: str
) -> ProductImport:
    """Parse `data`, upsert every valid row in chunks (phase 1), then embed
    whatever needs it (phase 2), and record the outcome on
    `product_imports`. An embedding failure in phase 2 does not fail the
    import -- see `_embed_imported_rows` and the module docstring.

    Commit discipline mirrors `app/rag/ingest.py::ingest_document`, for the
    same reason: `status=processing` is committed through its own
    independent transaction up front, so it is an observable state rather
    than one hidden inside a transaction that might run for minutes; and on
    any failure -- a structural parse error, a cancellation, or anything
    else escaping the chunk loop below -- `status=failed` is committed
    through a second, independent transaction rather than reusing whatever
    session was open when the failure happened. That session may itself be
    the thing that is broken (a database error inside a chunk's own
    `tenant_session` has already closed *that* transaction by the time
    control reaches here), so nothing below ever tries to keep using it.

    Each chunk gets its own `tenant_session` and its own try/except. A
    chunk that fails at the database is NOT simply reported as every one
    of its rows failing: Python-level validation (`ProductInput`, including its
    `Numeric(12, 2)`-mirroring `price` constraint) catches the overwhelming
    majority of bad input before a chunk is ever built, but it cannot catch
    every way a write can fail at the database -- a constraint this module
    never anticipated, a connection blip, a value that is valid `Decimal`
    but not valid for some column this module doesn't yet validate as
    tightly. Convicting all 500 rows in the chunk for one row's fault would
    be exactly the failure the brief exists to prevent, just moved one
    layer down from where the report first closed it. So a failing chunk is
    re-run through `_upsert_rows_individually`, one row per transaction,
    and only the row(s) that actually fail alone are reported failed --
    see that function's docstring. This fallback only runs on the error
    path: a clean import never pays a per-row round trip for it.
    """
    async with tenant_session(tenant) as session:
        await ProductImportService(session, tenant).mark_processing(product_import_id)

    try:
        parse_result = parse_import(data, mime_type)
        deduped_rows, duplicate_errors = _dedupe_by_external_id(parse_result.rows)
        errors: list[RowError] = [*parse_result.errors, *duplicate_errors]
        succeeded = 0
        failed = len(errors)
        pending_embedding: list[uuid.UUID] = []

        # Phase 1: product data, committed chunk by chunk.
        for chunk in _chunks(deduped_rows, _IMPORT_CHUNK_SIZE):
            try:
                pending_embedding.extend(await _upsert_chunk(tenant, chunk))
                succeeded += len(chunk)
            except DBAPIError as exc:
                logger.warning(
                    "product_import_chunk_failed_retrying_row_by_row",
                    product_import_id=str(product_import_id),
                    organization_id=str(tenant.organization_id),
                    rows=len(chunk),
                    error=bounded_error_message(exc),
                )
                row_pending, row_succeeded, row_errors = await _upsert_rows_individually(
                    tenant, chunk
                )
                pending_embedding.extend(row_pending)
                succeeded += row_succeeded
                errors.extend(row_errors)
                failed += len(row_errors)

        # Phase 2: vectors for whatever phase 1 said needs one.
        warning = await _embed_imported_rows(tenant, product_import_id, pending_embedding)
    except asyncio.CancelledError:
        async with tenant_session(tenant) as session:
            await ProductImportService(session, tenant).mark_failed(
                product_import_id,
                "import was cancelled before it finished -- most likely the worker's "
                "job timeout or a worker shutdown. Retrying is safe.",
            )
        raise
    except Exception as exc:
        async with tenant_session(tenant) as session:
            await ProductImportService(session, tenant).mark_failed(
                product_import_id, bounded_error_message(exc)
            )
        raise

    async with tenant_session(tenant) as session:
        return await ProductImportService(session, tenant).mark_completed(
            product_import_id,
            total_rows=parse_result.total_rows,
            succeeded_count=succeeded,
            failed_count=failed,
            errors=[error.to_dict() for error in errors],
            warning=warning,
        )
